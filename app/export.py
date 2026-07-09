"""Экспорт дашбордов в Excel через openpyxl (ТЗ §7).

Генерация в память (ФС Railway эфемерна). Учитываются активные фильтры.
Значения — без формул, с корректными форматами (проценты, суммы).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, metrics

# Форматы чисел Excel
F_INT = "#,##0"
F_MONEY = "#,##0 ₽"
F_PCT = "0.0%"
F_RATIO = "0.00"
F_DATE = "yyyy-mm-dd hh:mm"

_HEADER_FILL = PatternFill("solid", fgColor="1467F5")
_HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
_TOTAL_FONT = Font(bold=True)


def _write_sheet(ws, columns, rows, total=None):
    """columns: список (заголовок, ключ, формат). rows: список dict."""
    for ci, (title, _key, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(1, ci, title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    r = 2
    if total is not None:
        _write_row(ws, r, columns, total, bold=True)
        r += 1
    for row in rows:
        _write_row(ws, r, columns, row)
        r += 1

    # Ширины колонок и заморозка шапки
    for ci, (title, _key, _fmt) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = max(12, len(title) + 2)
    ws.freeze_panes = "A2"


def _write_row(ws, r, columns, data, bold=False):
    for ci, (_title, key, fmt) in enumerate(columns, start=1):
        val = data.get(key)
        cell = ws.cell(r, ci, val if val is not None else "—")
        if isinstance(val, (int, float)) and fmt:
            cell.number_format = fmt
        if bold:
            cell.font = _TOTAL_FONT


SRC_COLS = [
    ("Источник", "source", None), ("MQL", "mql", F_INT), ("SQL", "sql", F_INT),
    ("Назначено встреч", "meeting_scheduled", F_INT),
    ("Проведено встреч", "meeting_held", F_INT),
    ("Выставлено счетов", "invoiced", F_INT), ("Продажи", "sold", F_INT),
    ("Неквал", "unqualified", F_INT),
    ("MQL→SQL", "cr_mql_sql", F_PCT), ("Доходимость", "reachability", F_PCT),
    ("% неквала", "unqual_pct", F_PCT),
    ("Затраты", "amount", F_MONEY), ("CPL", "cpl", F_MONEY),
    ("Цена встречи", "price_meeting", F_MONEY), ("Цена продажи", "price_sale", F_MONEY),
    ("Выручка", "revenue", F_MONEY), ("ROAS", "roas", F_RATIO),
    ("Средний чек", "avg_check", F_MONEY),
]

MONTH_COLS = [
    ("Месяц", "month", None), ("MQL", "mql", F_INT), ("SQL", "sql", F_INT),
    ("Назначено встреч", "meeting_scheduled", F_INT),
    ("Проведено встреч", "meeting_held", F_INT),
    ("Выставлено счетов", "invoiced", F_INT), ("Продажи", "sold", F_INT),
    ("MQL→SQL", "cr_mql_sql", F_PCT), ("SQL→встреча", "cr_sql_scheduled", F_PCT),
    ("Доходимость", "cr_scheduled_held", F_PCT), ("Встреча→Счёт", "cr_held_invoiced", F_PCT),
    ("Счёт→Продажа", "cr_invoiced_sold", F_PCT), ("MQL→Продажа", "cr_mql_sale", F_PCT),
    ("% неквала", "unqual_pct", F_PCT), ("Выручка", "revenue", F_MONEY),
    ("Реклама", "ad_spend", F_MONEY), ("Затраты: маркетинг", "marketing_cost", F_MONEY),
    ("Затраты: продажи", "sales_cost", F_MONEY), ("Затраты ∑", "total_cost", F_MONEY),
    ("CAC", "cac", F_MONEY), ("Средний чек", "avg_check", F_MONEY),
    ("LTV", "ltv", F_MONEY), ("LTV:CAC", "ltv_cac", F_RATIO), ("ROAS", "roas", F_RATIO),
    ("Длина сделки, мес", "deal_length", F_RATIO),
]

MGR_COLS = [
    ("Менеджер", "manager", None), ("MQL", "mql", F_INT), ("SQL", "sql", F_INT),
    ("Назначено встреч", "meeting_scheduled", F_INT),
    ("Проведено встреч", "meeting_held", F_INT),
    ("Выставлено счетов", "invoiced", F_INT), ("Продажи", "sold", F_INT),
    ("MQL→SQL", "cr_mql_sql", F_PCT), ("Доходимость", "reachability", F_PCT),
    ("MQL→Продажа", "cr_mql_sale", F_PCT), ("% неквала", "unqual_pct", F_PCT),
    ("Выручка", "revenue", F_MONEY), ("Средний чек", "avg_check", F_MONEY),
]

DEAL_COLS = [
    ("ID", "amo_lead_id", None), ("Источник", "source", None),
    ("Менеджер", "manager", None), ("Статус", "status_name", None),
    ("Создана", "created_at", F_DATE), ("Закрыта", "closed_at", F_DATE),
    ("Сумма", "price", F_MONEY), ("Тип клиента", "client_type", None),
    ("Ежемес. платёж", "monthly_payment", F_MONEY),
    ("Этап (макс)", "max_stage_reached", F_INT),
    ("Продажа", "sold", None), ("Отказ", "is_lost", None),
    ("Источник этапа", "stage_source", None),
]

TOUCH_COLS = [
    ("Месяц", "month", None), ("Источник", "source", None),
    ("Сумма", "amount", F_MONEY), ("Касания", "touches", F_INT),
    ("Комментарий", "comment", None),
]

OPCOST_COLS = [
    ("Месяц", "month", None), ("Расходы на маркетинг", "cost_marketing", F_MONEY),
    ("Расходы на продажи", "cost_sales", F_MONEY), ("Комментарий", "comment", None),
]


def build_workbook(f: dict, config) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    src = metrics.by_source(f)
    _write_sheet(wb.create_sheet("По источникам"), SRC_COLS, src["rows"], src["total"])

    mon = metrics.by_month(f, config)
    _write_sheet(wb.create_sheet("По месяцам"), MONTH_COLS, mon["rows"], mon["total"])

    mgr = metrics.by_manager(f)
    _write_sheet(wb.create_sheet("По менеджерам"), MGR_COLS, mgr["rows"], mgr["total"])

    _write_sheet(wb.create_sheet("Сделки"), DEAL_COLS, _deals_rows(f))
    _write_sheet(wb.create_sheet("Реклама (затраты)"), TOUCH_COLS, _touch_rows())
    _write_sheet(wb.create_sheet("Затраты по месяцам"), OPCOST_COLS, _opcost_rows())

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _deals_rows(f: dict) -> list:
    where, params = metrics._deal_where(f)
    rows = db.query(
        f"""SELECT d.amo_lead_id, d.source,
                   COALESCE(m.name, '') AS manager,
                   COALESCE(s.name, '') AS status_name,
                   d.created_at, d.closed_at, d.price, d.client_type,
                   d.monthly_payment, d.max_stage_reached,
                   CASE WHEN d.sold THEN 'да' ELSE '' END AS sold,
                   CASE WHEN d.is_lost THEN 'да' ELSE '' END AS is_lost,
                   d.stage_source
              FROM deals d
              LEFT JOIN managers m ON m.amo_user_id = d.responsible_user_id
              LEFT JOIN pipeline_statuses s ON s.status_id = d.current_status_id
             WHERE {where}
             ORDER BY d.created_at DESC""",
        params,
    )
    # created_at/closed_at приходят с tz — openpyxl не поддерживает tz-aware datetime
    out = []
    for r in rows:
        r = dict(r)
        for k in ("created_at", "closed_at"):
            if r[k] is not None:
                r[k] = r[k].replace(tzinfo=None)
        out.append(r)
    return out


def _touch_rows() -> list:
    return db.query(
        """SELECT to_char(period_month, 'YYYY-MM') AS month, source, amount,
                  touches, comment FROM costs_touches
             ORDER BY period_month DESC, source"""
    )


def _opcost_rows() -> list:
    return db.query(
        """SELECT to_char(period_month, 'YYYY-MM') AS month, cost_marketing,
                  cost_sales, comment FROM monthly_costs ORDER BY period_month DESC"""
    )
